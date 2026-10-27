"""Join INSEE COG 2026 and the 2025 density grid (2026 commune geography)."""
import argparse
import csv
import json
from pathlib import Path

from openpyxl import load_workbook


def read_csv(file):
    with file.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def build(commune_file: Path, department_file: Path, density_file: Path, output: Path):
    departments = {row["DEP"]: row["NCCENR"] for row in read_csv(department_file)}
    book = load_workbook(density_file, read_only=True, data_only=True)
    sheet = book["Maille communale"]
    density = {}
    for row in sheet.iter_rows(min_row=6, values_only=True):
        code = str(row[0] or "").zfill(5)
        if len(code) == 5 and str(row[2]) in {"1", "2", "3"}:
            density[code] = {
                "density": {"1": "dense", "2": "intermediate", "3": "rural"}[str(row[2])],
                "population2023": int(row[4]) if isinstance(row[4], (float, int)) else None,
            }
    records = []
    for row in read_csv(commune_file):
        if row["TYPECOM"] not in {"COM", "ARM"}:
            continue
        code = row["COM"]
        records.append({
            "communeCode": code, "commune": row["NCCENR"], "communeLabel": row["LIBELLE"],
            "departmentCode": row["DEP"], "department": departments.get(row["DEP"], ""),
            "density": density.get(code, {}).get("density"),
            "population2023": density.get(code, {}).get("population2023"), "sourceYear": 2026,
            "parentCommuneCode": row["COMPARENT"] or None,
        })
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps({
        "sourceCog": "https://www.insee.fr/fr/information/8740222",
        "sourceDensity": "https://www.insee.fr/fr/information/8571524",
        "records": records,
    }, ensure_ascii=False), encoding="utf8")
    print(json.dumps({"records": len(records), "withDensity": sum(bool(r["density"]) for r in records), "output": str(output)}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--communes", type=Path, required=True)
    parser.add_argument("--departments", type=Path, required=True)
    parser.add_argument("--density", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    build(args.communes, args.departments, args.density, args.out)
